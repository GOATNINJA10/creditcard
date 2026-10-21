from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = ROOT / "data" / "source" / "twitterfiles.csv"
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "outputs"
SOURCE_URL = (
    "https://github.com/kyasanur/"
    "Sentiment-Analysis-on-Credit-Cards-using-Online-Reviews"
)

RELEVANCE_TERMS = (
    r"credit card|creditcard|credit-card|\bcard\b|cashback|cash back|"
    r"annual fee|joining fee|lounge|reward|points|statement|emi|upi|"
    r"forex|surcharge"
)
ISSUER_NAMES = {
    "SBI": "SBI Card",
    "hdfc": "HDFC Bank",
    "Axis": "Axis Bank",
    "icici": "ICICI Bank",
}


def sentiment_label(score: float) -> str:
    if score > 0:
        return "Positive"
    if score < 0:
        return "Negative"
    return "Neutral"


def prepare_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(f"Twitter source file not found: {SOURCE_PATH}")

    raw = pd.read_csv(SOURCE_PATH)
    required = {
        "ttext",
        "date",
        "retweetcount",
        "favoritecount",
        "score",
        "product",
        "country",
    }
    missing = sorted(required - set(raw.columns))
    if missing:
        raise ValueError(f"Twitter source is missing columns: {', '.join(missing)}")

    raw_count = len(raw)
    text = raw["ttext"].fillna("").astype(str).str.replace(r"\s+", " ", regex=True).str.strip()
    eligible = raw.assign(review_text=text)
    eligible = eligible[
        eligible["country"].fillna("").eq("India")
        & eligible["product"].isin(ISSUER_NAMES)
        & eligible["review_text"].str.contains(RELEVANCE_TERMS, case=False, regex=True)
        & eligible["review_text"].ne("")
    ].copy()
    eligible = eligible.drop_duplicates(subset=["review_text"]).reset_index(drop=True)

    result = pd.DataFrame(
        {
            "record_id": [f"TW{i:04d}" for i in range(1, len(eligible) + 1)],
            "bank_name": eligible["product"].map(ISSUER_NAMES),
            "card_name": "Explicit card name not consistently available",
            "source": "Published GitHub Twitter dataset",
            "record_type": "twitter_mention",
            "published_date": pd.to_datetime(eligible["date"], dayfirst=True, errors="coerce"),
            "review_text": eligible["review_text"],
            "retweet_count": pd.to_numeric(eligible["retweetcount"], errors="coerce").fillna(0),
            "favorite_count": pd.to_numeric(eligible["favoritecount"], errors="coerce").fillna(0),
            "engagement": (
                pd.to_numeric(eligible["retweetcount"], errors="coerce").fillna(0)
                + pd.to_numeric(eligible["favoritecount"], errors="coerce").fillna(0)
            ),
            "source_url": SOURCE_URL,
            "source_sentiment_score": pd.to_numeric(eligible["score"], errors="coerce").fillna(0),
        }
    )
    result["sentiment"] = result["source_sentiment_score"].map(sentiment_label)
    result["month"] = result["published_date"].dt.to_period("M").astype(str)
    result["input_records"] = raw_count
    result.attrs["india_records"] = int(raw["country"].fillna("").eq("India").sum())
    return raw, result


def build_summaries(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    mentions = (
        df.groupby("bank_name")
        .size()
        .reset_index(name="mentions")
        .sort_values("mentions", ascending=False)
    )
    sentiment = (
        pd.crosstab(df["bank_name"], df["sentiment"], normalize="index")
        .mul(100)
        .round(1)
        .reindex(columns=["Positive", "Neutral", "Negative"], fill_value=0)
        .reset_index()
    )
    monthly = (
        df.groupby(["month", "bank_name"])
        .size()
        .reset_index(name="mentions")
        .pivot(index="month", columns="bank_name", values="mentions")
        .fillna(0)
        .reset_index()
    )
    engagement = (
        df.groupby("bank_name")
        .agg(
            mentions=("record_id", "size"),
            total_retweets=("retweet_count", "sum"),
            total_favourites=("favorite_count", "sum"),
            average_engagement=("engagement", "mean"),
        )
        .reset_index()
        .sort_values("mentions", ascending=False)
    )
    quality = pd.DataFrame(
        {
            "check": [
                "Raw source records",
                "Records from India",
                "Relevant unique Twitter mentions",
                "Issuers covered",
                "Date coverage (%)",
                "Duplicate/irrelevant rows excluded",
            ],
            "value": [
                int(df["input_records"].iloc[0]),
                int(df.attrs["india_records"]),
                len(df),
                df["bank_name"].nunique(),
                round(df["published_date"].notna().mean() * 100, 1),
                int(df["input_records"].iloc[0] - len(df)),
            ],
        }
    )
    return {
        "Mentions": mentions,
        "Sentiment": sentiment,
        "Monthly_Mentions": monthly,
        "Engagement": engagement,
        "Quality": quality,
    }


def build_outputs(df: pd.DataFrame, summaries: dict[str, pd.DataFrame]) -> None:
    DATA_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)
    df.to_csv(DATA_DIR / "twitter_listening_data.csv", index=False, date_format="%Y-%m-%d")

    workbook_path = OUTPUT_DIR / "Indian_Credit_Card_Social_Listening_PoW.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Data", index=False)
        for name, table in summaries.items():
            table.to_excel(writer, sheet_name=name, index=False)
        pd.DataFrame({"Dashboard note": ["Historical Twitter listening module; not customer-review data."]}).to_excel(
            writer, sheet_name="Dashboard", index=False
        )

    wb = load_workbook(workbook_path)
    dashboard = wb["Dashboard"]
    dashboard.delete_rows(1, dashboard.max_row)
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:N1")
    dashboard["A1"] = "HISTORICAL TWITTER CREDIT-CARD LISTENING"
    dashboard["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    dashboard["A1"].alignment = Alignment(horizontal="center")
    dashboard.merge_cells("A2:N2")
    dashboard["A2"] = (
        "Filtered Indian Twitter mentions | 19 January–19 March 2022 | "
        "Separate from the pure customer-review analysis"
    )
    dashboard["A2"].font = Font(italic=True, color="666666")
    dashboard["A2"].alignment = Alignment(horizontal="center")
    dashboard.merge_cells("A4:N4")
    dashboard["A4"] = (
        "Interpretation note: mention volume is a source-volume proxy, not market share. "
        "Tweets are social-listening mentions, not necessarily customer reviews."
    )
    dashboard["A4"].font = Font(italic=True, color="C00000")
    dashboard["A4"].alignment = Alignment(wrap_text=True)

    mentions_sheet = wb["Mentions"]
    chart = BarChart()
    chart.title = "Twitter mentions by issuer"
    chart.x_axis.title = "Issuer"
    chart.y_axis.title = "Number of mentions"
    chart.add_data(
        Reference(mentions_sheet, min_col=2, min_row=1, max_row=mentions_sheet.max_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(mentions_sheet, min_col=1, min_row=2, max_row=mentions_sheet.max_row)
    )
    chart.height = 7
    chart.width = 13
    dashboard.add_chart(chart, "A6")

    monthly_sheet = wb["Monthly_Mentions"]
    line = LineChart()
    line.title = "Monthly Twitter mentions"
    line.x_axis.title = "Month"
    line.y_axis.title = "Number of mentions"
    line.add_data(
        Reference(
            monthly_sheet,
            min_col=2,
            max_col=monthly_sheet.max_column,
            min_row=1,
            max_row=monthly_sheet.max_row,
        ),
        titles_from_data=True,
    )
    line.set_categories(
        Reference(monthly_sheet, min_col=1, min_row=2, max_row=monthly_sheet.max_row)
    )
    line.height = 7
    line.width = 18
    dashboard.add_chart(line, "H6")

    engagement_sheet = wb["Engagement"]
    engagement_chart = BarChart()
    engagement_chart.title = "Average engagement by issuer"
    engagement_chart.x_axis.title = "Issuer"
    engagement_chart.y_axis.title = "Average retweets + favourites"
    engagement_chart.add_data(
        Reference(
            engagement_sheet,
            min_col=5,
            min_row=1,
            max_row=engagement_sheet.max_row,
        ),
        titles_from_data=True,
    )
    engagement_chart.set_categories(
        Reference(engagement_sheet, min_col=1, min_row=2, max_row=engagement_sheet.max_row)
    )
    engagement_chart.height = 7
    engagement_chart.width = 13
    dashboard.add_chart(engagement_chart, "A22")

    for column in range(1, 15):
        dashboard.column_dimensions[chr(64 + column)].width = 14
    wb.save(workbook_path)

    chart_dir = OUTPUT_DIR / "charts"
    chart_dir.mkdir(exist_ok=True)
    summaries["Mentions"].plot.bar(
        x="bank_name", y="mentions", legend=False, title="Historical Twitter mentions by issuer"
    ).set(xlabel="Issuer", ylabel="Number of mentions")
    plt.tight_layout()
    plt.savefig(chart_dir / "twitter_mentions.png", dpi=160)
    plt.close()

    monthly_plot = summaries["Monthly_Mentions"].set_index("month")
    monthly_plot.plot(kind="line", marker="o", title="Monthly Twitter mentions")
    plt.xlabel("Month")
    plt.ylabel("Number of mentions")
    plt.tight_layout()
    plt.savefig(chart_dir / "twitter_monthly_mentions.png", dpi=160)
    plt.close()

    engagement_plot = summaries["Engagement"]
    engagement_plot.plot.bar(
        x="bank_name",
        y="average_engagement",
        legend=False,
        title="Average Twitter engagement by issuer",
    ).set(xlabel="Issuer", ylabel="Average retweets + favourites")
    plt.tight_layout()
    plt.savefig(chart_dir / "twitter_engagement.png", dpi=160)
    plt.close()


def main() -> None:
    _, df = prepare_data()
    summaries = build_summaries(df)
    build_outputs(df, summaries)
    print(f"Created {len(df)} filtered unique Twitter mentions")
    print(df["bank_name"].value_counts().to_string())


if __name__ == "__main__":
    main()
